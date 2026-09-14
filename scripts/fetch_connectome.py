"""Download Cook et al. 2019 connectome tables and extract neuron-to-neuron edges.

Run:  uv run --with openpyxl scripts/fetch_connectome.py [--out data/connectome/cook2019]

Writes neurons.csv, chemical.csv, gap_junction.csv. Only the 302 hermaphrodite
neurons are kept; muscles, glia and other end organs are dropped.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import tempfile
import urllib.request
from pathlib import Path

import openpyxl

BASE = "https://wormwiring.org/si/"
SI5 = "SI 5 Connectome adjacency matrices, corrected July 2020.xlsx"
SI4 = "SI 4 Cell lists.xlsx"
NEURON_TYPES = {"sensory", "interneuron", "motorneuron", "neuron"}


def download(name: str, into: Path) -> Path:
    target = into / name
    url = BASE + urllib.request.quote(name)
    print(f"downloading {url}")
    urllib.request.urlretrieve(url, target)
    return target


def read_cells(si4: Path) -> dict[str, tuple[str, str, str]]:
    """name -> (type, category, notes) for pharynx, sex-shared and hermaphrodite cells."""
    wb = openpyxl.load_workbook(si4, read_only=True)
    cells: dict[str, tuple[str, str, str]] = {}
    for sheet in ("pharynx", "sex-shared", "hermaphrodite specific"):
        for row in wb[sheet].iter_rows(values_only=True):
            name, ctype = row[0], row[1]
            if not name or not ctype or name == "name":
                continue
            category = str(row[2]) if len(row) > 2 and row[2] else ""
            notes = str(row[3]) if len(row) > 3 and row[3] else ""
            cells[str(name)] = (str(ctype), category, notes)
    return cells


def read_matrix(si5: Path, sheet: str) -> dict[tuple[str, str], int]:
    """(row cell, column cell) -> weight for every non-zero entry."""
    wb = openpyxl.load_workbook(si5, read_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    cols = rows[2][3:]
    out: dict[tuple[str, str], int] = {}
    for row in rows[3:]:
        src = row[2]
        if src is None:
            continue
        for col, value in zip(cols, row[3:], strict=False):
            if col is not None and value:
                out[(str(src), str(col))] = int(value)
    return out


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=Path("data/connectome/cook2019"))
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as tmp:
        si5 = download(SI5, Path(tmp))
        si4 = download(SI4, Path(tmp))
        cells = read_cells(si4)
        chemical = read_matrix(si5, "hermaphrodite chemical")
        gap = read_matrix(si5, "hermaphrodite gap jn symmetric")

    neurons = sorted(n for n, (t, _, _) in cells.items() if t in NEURON_TYPES)
    nset = set(neurons)
    print(f"neurons: {len(neurons)}")

    with (args.out / "neurons.csv").open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["name", "type", "category", "notes"])
        for n in neurons:
            t, cat, notes = cells[n]
            w.writerow([n, t, cat, notes])

    chem_edges = sorted((s, t, v) for (s, t), v in chemical.items() if s in nset and t in nset)
    with (args.out / "chemical.csv").open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["source", "target", "weight"])
        w.writerows(chem_edges)
    print(f"chemical edges: {len(chem_edges)}")

    pairs: dict[tuple[str, str], int] = {}
    for (a, b), v in gap.items():
        if a in nset and b in nset:
            key = (a, b) if a <= b else (b, a)
            if key in pairs and pairs[key] != v:
                raise SystemExit(f"asymmetric gap junction {key}: {pairs[key]} vs {v}")
            pairs[key] = v
    gap_edges = sorted((a, b, v) for (a, b), v in pairs.items())
    with (args.out / "gap_junction.csv").open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["a", "b", "weight"])
        w.writerows(gap_edges)
    print(f"gap junction pairs: {len(gap_edges)}")

    (args.out / "RETRIEVED").write_text(
        f"{dt.date.today().isoformat()}\n{BASE}{SI5}\n{BASE}{SI4}\n"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
